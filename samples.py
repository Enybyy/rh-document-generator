"""Synthetic examples; no employee data."""
from io import BytesIO
from docx import Document
from docx.shared import Pt, Cm
from openpyxl import Workbook

ROWS = [
    ["00000001", "Valeria Campos Medina", "Analista administrativa", "Lima", "01/10/2026", "31/12/2026", 2800],
    ["00000002", "Mateo Salazar Torres", "Asistente de operaciones", "Arequipa", "01/10/2026", "31/12/2026", 2400],
    ["00000003", "Lucía Paredes Ríos", "Coordinadora de proyecto", "Cusco", "01/10/2026", "31/12/2026", 3600],
]
HEADERS = ["DNI", "NOMBRE", "CARGO", "CIUDAD", "FECHA INICIO", "FECHA FIN", "PAGO"]


def excel():
    book = Workbook()
    sheet = book.active
    sheet.title = "Personal de muestra"
    sheet.append(HEADERS)
    for row in ROWS:
        sheet.append(row)
    for col in "ABCDEFG":
        sheet.column_dimensions[col].width = 29 if col != "B" else 36
    sheet.freeze_panes = "A2"
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


def template():
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.5)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading("Ficha de servicios", 0)
    doc.add_paragraph("Ejemplo de automatización documental con datos ficticios. Documento de muestra para revisión administrativa.")
    doc.add_heading("Datos del colaborador", 1)
    table = doc.add_table(rows=0, cols=2)
    table.style = "Light Shading Accent 1"
    for label, field in [("Nombre", "NOMBRE"), ("DNI de muestra", "DNI"), ("Servicio", "CARGO"), ("Ciudad", "CIUDAD"), ("Inicio", "FECHA_INICIO"), ("Fin", "FECHA_FIN"), ("Pago S/", "PAGO")]:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = label, f"[{field}]"
    doc.add_paragraph("Observaciones: ______________________________________________")
    section.footer.paragraphs[0].text = "Generar RH · Documento de muestra"
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()
