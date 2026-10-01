from io import BytesIO
from zipfile import ZipFile
import json
import pytest
from docx import Document
from openpyxl import Workbook, load_workbook

from generator import InputError, prepare, build_zip, read_excel, load_template
from app import app
import samples


def workbook(headers, rows):
    book = Workbook()
    sheet = book.active
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


def template_doc():
    doc = Document()
    p = doc.add_paragraph("Para ")
    p.add_run("[NOM").bold = True
    p.add_run("BRE]").italic = True
    p.add_run(" / [DNI] / [NOMBRE]")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "{{ CARGO }}"
    nested = table.cell(0, 0).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "[CIUDAD]"
    doc.sections[0].header.paragraphs[0].text = "[NOMBRE]"
    doc.sections[0].footer.paragraphs[0].text = "[DNI]"
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def test_complete_generation_split_format_tables_headers_footers():
    t = template_doc()
    data = prepare(samples.excel(), t)
    out = build_zip(t, data, {0, 2})
    with ZipFile(BytesIO(out)) as z:
        names = [n for n in z.namelist() if n.endswith('.docx')]
        assert len(names) == 2
        doc = Document(BytesIO(z.read(names[0])))
        assert doc.paragraphs[0].text == 'Para Valeria Campos Medina / 00000001 / Valeria Campos Medina'
        assert doc.paragraphs[0].runs[1].bold is True
        assert doc.tables[0].cell(0, 0).paragraphs[0].text == 'Analista administrativa'
        assert doc.tables[0].cell(0, 0).tables[0].cell(0, 0).text == 'Lima'
        assert doc.sections[0].header.paragraphs[0].text == 'Valeria Campos Medina'
        assert doc.sections[0].footer.paragraphs[0].text == '00000001'
        sheet = load_workbook(BytesIO(z.read('revision.xlsx'))).active
        assert sheet.cell(3, 4).value == 'No seleccionado'
        assert json.loads(z.read('resumen.json'))['generados'] == 2


def test_duplicate_names_do_not_overwrite():
    values = list(samples.ROWS[0])
    source = workbook(samples.HEADERS, [values, values])
    result = prepare(source, samples.template())
    with ZipFile(BytesIO(build_zip(samples.template(), result, {0, 1}))) as z:
        assert len([n for n in z.namelist() if n.endswith('.docx')]) == 2


def test_numeric_dni_keeps_leading_zeros():
    assert read_excel(workbook(['DNI'], [[12345]]))[0]['DNI'] == '00012345'


@pytest.mark.parametrize('headers,rows', [(['DNI', 'DNI'], [['1', '2']]), (['DNI', None], [['1', '2']]), (['DNI'], []), (['DNI'], [['=1+1']]), (['DNI'], [['1']]*501)])
def test_bad_workbooks_rejected(headers, rows):
    with pytest.raises(InputError):
        read_excel(workbook(headers, rows))


@pytest.mark.parametrize('field,value', [('DNI','123'), ('DNI','１２３４５６７８'), ('PAGO',-1), ('PAGO','nan'), ('PAGO','2800,50'), ('PAGO','1,2,3'), ('PAGO','1.234'), ('FECHA INICIO','31/02/2026'), ('FECHA FIN','01/01/2025'), ('NOMBRE',None)])
def test_invalid_rows_excluded(field, value):
    values = list(samples.ROWS[0])
    values[samples.HEADERS.index(field)] = value
    data = prepare(workbook(samples.HEADERS, [values]), samples.template())
    assert data['rows'][0]['valid'] is False
    with pytest.raises(InputError):
        build_zip(samples.template(), data, {0})


def test_personal_join_by_dni_and_missing():
    p = workbook(['DNI', 'NOMBRE'], [['00000001', 'Nombre de referencia']])
    data = prepare(samples.excel(), samples.template(), p)
    assert data['rows'][0]['values']['NOMBRE'] == 'Nombre de referencia'
    assert data['rows'][0]['valid']
    assert not data['rows'][1]['valid']


def test_duplicate_personal_blocks_ambiguous_join():
    with pytest.raises(InputError, match='duplicado'):
        prepare(samples.excel(), samples.template(), workbook(['DNI'], [['00000001'], ['00000001']]))


def test_unknown_template_tag_reports_missing_column():
    doc = Document()
    doc.add_paragraph('[DNI] [CAMPO_AUSENTE]')
    stream = BytesIO()
    doc.save(stream)
    data = prepare(samples.excel(), stream.getvalue())
    assert all(not r['valid'] for r in data['rows'])


def test_paths_in_names_are_sanitized_and_formula_report_escaped():
    values = list(samples.ROWS[0])
    values[1] = '+../../../Person'
    result = prepare(workbook(samples.HEADERS, [values]), samples.template())
    with ZipFile(BytesIO(build_zip(samples.template(), result, {0}))) as z:
        assert all('..' not in n for n in z.namelist())
        sheet = load_workbook(BytesIO(z.read('revision.xlsx'))).active
        assert sheet.cell(2, 3).data_type != 'f'


@pytest.mark.parametrize('data', [b'not a docx', samples.excel()])
def test_fake_template_rejected(data):
    with pytest.raises(InputError):
        load_template(data)


def upload_data():
    return {'excel': (BytesIO(samples.excel()), 'personal.xlsx'), 'template': (BytesIO(samples.template()), 'plantilla.docx')}


def test_api_preview_download_and_invalid_selection():
    client = app.test_client()
    r = client.post('/api/preview', data=upload_data())
    assert r.status_code == 200 and len(r.json['rows']) == 3
    data = upload_data()
    data['selected'] = '[0, 1]'
    r = client.post('/api/generate', data=data)
    assert r.status_code == 200 and r.mimetype == 'application/zip'
    with ZipFile(BytesIO(r.data)) as z:
        assert len([n for n in z.namelist() if n.endswith('.docx')]) == 2
    for value in ('null', '[true]', '[-1]', '[999]', 'oops'):
        data = upload_data()
        data['selected'] = value
        assert client.post('/api/generate', data=data).status_code == 400
    assert client.post('/api/preview', data={}).status_code == 400
    assert client.get('/api/samples/unknown').status_code == 404


def test_independent_requests_and_no_disk_uploads():
    first, second = app.test_client(), app.test_client()
    assert first.post('/api/preview', data=upload_data()).status_code == 200
    assert second.post('/api/generate', data={'selected': '[0]'}).status_code == 400
    assert first.get('/api/health').headers['Cache-Control'] == 'no-store'


def test_large_output_blocked_before_generation(monkeypatch):
    monkeypatch.setattr('generator.MAX_OUTPUT_BYTES', 30)
    data = prepare(samples.excel(), samples.template())
    with pytest.raises(InputError, match='64 MB'):
        build_zip(samples.template(), data, {0, 1, 2})
