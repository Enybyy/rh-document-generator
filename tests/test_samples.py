"""The shipped contract must remain compatible with the real generation path."""
from io import BytesIO
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn

import samples
from generator import prepare, build_zip


def test_reference_contract_has_two_explicit_pages_and_editable_fields():
    template = samples.template()
    preview = prepare(samples.excel(), template)
    assert set(preview["fields"]) == {
        "NOMBRE", "DNI", "CIUDAD", "CARGO", "FECHA_INICIO", "FECHA_FIN", "PAGO"
    }
    with ZipFile(BytesIO(build_zip(template, preview, {0}))) as bundle:
        filename = next(n for n in bundle.namelist() if n.endswith(".docx"))
        document = Document(BytesIO(bundle.read(filename)))
    text = "\n".join(p.text for p in document.paragraphs)
    assert "Contrato de prestación de servicios" in text
    assert "Términos del servicio" in text
    assert "Valeria Campos Medina" in text
    assert "S/ 2800.00" in text
    assert "06 Conformidad y firmas" in text
    assert "[NOMBRE]" not in text
    breaks = document.element.xpath('//w:br[@w:type="page"]')
    assert len(breaks) == 1
    section = document.sections[0]
    assert round(section.page_width.cm, 1) == 21.0
    assert round(section.page_height.cm, 1) == 29.7
    assert section.header.paragraphs[0].text.endswith("Valeria Campos Medina")
    fields = section.footer._element.xpath(".//w:fldSimple")
    assert any(field.get(qn("w:instr")) == "PAGE" for field in fields)
