"""Excel + Word mail merge. All processing stays in the current request."""
from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
import json
import re
import unicodedata
from zipfile import ZipFile, ZIP_DEFLATED, BadZipFile

from docx import Document
from docx.oxml.ns import qn
from openpyxl import Workbook, load_workbook

MAX_ROWS = 500
MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_OUTPUT_BYTES = 64 * 1024 * 1024
TAG = re.compile(r"\[([^\[\]\r\n]+)\]|\{\{\s*([^{}\r\n]+?)\s*\}\}")


class InputError(ValueError):
    """A recoverable error safe to present to the user."""


def normalize(value):
    text = unicodedata.normalize("NFKD", str(value or "").strip())
    text = "".join(c for c in text if not unicodedata.combining(c))
    return re.sub(r"[^A-Z0-9]+", "_", text.upper()).strip("_")


ALIASES = {
    "NUMERO_DE_DOCUMENTO": "DNI", "NUMERO_DOCUMENTO": "DNI",
    "NRO_DE_DOCUMENTO": "DNI", "NRO_DOCUMENTO": "DNI",
    "APELLIDOS_Y_NOMBRES": "NOMBRE", "NOMBRE_COMPLETO": "NOMBRE",
    "APELLIDOS_Y_NOMBRES_VERIFICADOS": "NOMBRE",
    "SALARIO": "PAGO", "FECHA_DE_INICIO": "FECHA_INICIO",
    "FECHA_DE_FIN": "FECHA_FIN", "UBICACION": "DIRECCION",
    "TIPO_DE_DOCUMENTO": "TIPO_DOCUMENTO",
}


def key(value):
    n = normalize(value)
    return ALIASES.get(n, n)


def safe_package(data, extension):
    if not data or len(data) > MAX_FILE_BYTES:
        raise InputError("Cada archivo debe contener datos y pesar como máximo 10 MB.")
    try:
        with ZipFile(BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > 3000 or sum(i.file_size for i in entries) > 80 * 1024 * 1024:
                raise InputError("El archivo descomprimido excede el límite permitido.")
            if any(i.flag_bits & 1 for i in entries):
                raise InputError("No se admiten archivos cifrados.")
            expected = "word/document.xml" if extension == "docx" else "xl/workbook.xml"
            if expected not in archive.namelist():
                raise InputError(f"El archivo no contiene un documento {extension.upper()} válido.")
            if any("vbaProject" in i.filename for i in entries):
                raise InputError("No se admiten documentos con macros.")
    except BadZipFile as exc:
        raise InputError(f"No se pudo abrir el archivo {extension.upper()}.") from exc


def read_excel(data):
    safe_package(data, "xlsx")
    try:
        book = load_workbook(BytesIO(data), read_only=True, data_only=False)
        try:
            sheet = book.worksheets[0]
            # Ignore malicious worksheet dimension hints; stream actual XML rows.
            sheet.reset_dimensions()
            iterator = sheet.iter_rows()
            first = next(iterator, ())
            if len(first) > 100:
                raise InputError("Se admiten como máximo 100 columnas.")
            headers = [key(c.value) for c in first]
            if not headers or any(not h for h in headers) or len(set(headers)) != len(headers):
                raise InputError("La primera fila debe tener encabezados únicos y sin celdas vacías.")
            rows = []
            for number, cells in enumerate(iterator, 2):
                if number > MAX_ROWS + 1:
                    raise InputError(f"Se admiten como máximo {MAX_ROWS} filas de datos.")
                if len(cells) > len(headers) and any(c.value is not None for c in cells[len(headers):]):
                    raise InputError(f"Fila {number}: hay datos sin encabezado.")
                if all(c.value is None for c in cells):
                    continue
                if any(c.data_type == "f" for c in cells):
                    raise InputError(f"Fila {number}: sustituye las fórmulas por valores antes de cargar.")
                row = {h: (cells[i].value if i < len(cells) else None) for i, h in enumerate(headers)}
                if isinstance(row.get("DNI"), (int, float)) and not isinstance(row.get("DNI"), bool):
                    numeric = row["DNI"]
                    row["DNI"] = f"{int(numeric):08d}" if numeric == int(numeric) else str(numeric)
                row["_fila"] = number
                rows.append(row)
            if not rows:
                raise InputError("El Excel no tiene filas de datos.")
            return rows
        finally:
            book.close()
    except InputError:
        raise
    except Exception as exc:
        raise InputError("No se pudo leer el Excel. Usa un archivo XLSX sin protección.") from exc


def paragraphs(document):
    """Include nested tables, headers and footers, without flattening formatting."""
    roots = [document.element.body]
    for section in document.sections:
        for name in ("header", "first_page_header", "even_page_header",
                     "footer", "first_page_footer", "even_page_footer"):
            part = getattr(section, name)
            if not part.is_linked_to_previous:
                roots.append(part._element)
    seen = set()
    for root in roots:
        for p in root.iter(qn("w:p")):
            if p not in seen:
                seen.add(p)
                yield p


def text_nodes(paragraph):
    # Textboxes may nest another paragraph inside a paragraph. Treat independently.
    return [t for t in paragraph.iter(qn("w:t"))
            if next(t.iterancestors(qn("w:p")), None) is paragraph]


def text_of(paragraph):
    return "".join(t.text or "" for t in text_nodes(paragraph))


def load_template(data):
    safe_package(data, "docx")
    try:
        doc = Document(BytesIO(data))
        tags = {key(m.group(1) or m.group(2)) for p in paragraphs(doc) for m in TAG.finditer(text_of(p))}
        if not tags:
            raise InputError("La plantilla no tiene etiquetas. Usa [NOMBRE] o {{ NOMBRE }}.")
        return doc, tags
    except InputError:
        raise
    except Exception as exc:
        raise InputError("No se pudo leer la plantilla Word.") from exc


def replace_paragraph(p, values):
    nodes = text_nodes(p)
    original = "".join(t.text or "" for t in nodes)
    starts, offset = [], 0
    for t in nodes:
        starts.append(offset)
        offset += len(t.text or "")
    for match in reversed(list(TAG.finditer(original))):
        start, end = match.span()
        left = next(i for i, t in enumerate(nodes) if starts[i] <= start < starts[i] + len(t.text or ""))
        right = next(i for i, t in enumerate(nodes) if starts[i] < end <= starts[i] + len(t.text or ""))
        replacement = values[key(match.group(1) or match.group(2))]
        prefix = (nodes[left].text or "")[:start - starts[left]]
        suffix = (nodes[right].text or "")[end - starts[right]:]
        nodes[left].text = prefix + replacement + (suffix if left == right else "")
        nodes[left].set(qn("xml:space"), "preserve")
        if left != right:
            for i in range(left + 1, right):
                nodes[i].text = ""
            nodes[right].text = suffix


def format_value(field, value):
    if value is None:
        return ""
    if field.startswith("FECHA"):
        if isinstance(value, (date, datetime)):
            return value.strftime("%d/%m/%Y")
        text = str(value).strip()
        for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
            try:
                return datetime.strptime(text, fmt).strftime("%d/%m/%Y")
            except ValueError:
                pass
        raise InputError(f"{field}: usa una fecha válida como 01/10/2026.")
    if field in {"PAGO", "TOTAL"}:
        try:
            raw = str(value).strip()
            if not re.fullmatch(r'[0-9]+(?:\.[0-9]{1,2})?', raw):
                raise InvalidOperation
            amount = Decimal(raw)
            if not amount.is_finite() or amount < 0 or amount > Decimal('999999999999.99'):
                raise InvalidOperation
            return f"{amount:.2f}"
        except InvalidOperation as exc:
            raise InputError(f"{field}: usa punto decimal y hasta dos decimales, sin separadores de miles; por ejemplo 2500.00.") from exc
    text = str(value).strip()
    if any(ord(c) < 32 and c not in '\t\n\r' for c in text):
        raise InputError(f"{field}: contiene caracteres de control no admitidos en Word.")
    return text


def prepare(excel, template, personal=None):
    _, tags = load_template(template)
    source = read_excel(excel)
    people = {}
    if personal:
        for row in read_excel(personal):
            dni = str(row.get("DNI", "")).strip()
            if not re.fullmatch(r"[0-9]{8}", dni):
                raise InputError(f"Personal, fila {row['_fila']}: DNI inválido.")
            if dni in people:
                raise InputError(f"Personal, fila {row['_fila']}: DNI duplicado; corrige el cruce ambiguo.")
            people[dni] = row
    results = []
    for index, row in enumerate(source):
        errors = []
        dni = str(row.get("DNI", "")).strip()
        if not re.fullmatch(r"[0-9]{8}", dni):
            errors.append("DNI: se requieren 8 dígitos.")
        if personal:
            if dni not in people:
                errors.append("No se encontró el DNI en la base de personal.")
            else:
                row = {**row, **{k: v for k, v in people[dni].items() if k != "_fila" and v is not None}}
        if "APELLIDO" in tags:
            pass  # Legacy templates can keep NOMBRE and APELLIDO separately.
        elif row.get("APELLIDO") and row.get("NOMBRE"):
            row["NOMBRE"] = f"{row['NOMBRE']} {row['APELLIDO']}"
        values = {}
        for field, value in row.items():
            if field == "_fila":
                continue
            try:
                values[field] = format_value(field, value)
            except InputError as exc:
                errors.append(str(exc))
                values[field] = ""
        for field in sorted(tags):
            if not values.get(field):
                errors.append(f"Falta el valor de {field}.")
        if values.get("FECHA_INICIO") and values.get("FECHA_FIN"):
            if datetime.strptime(values["FECHA_FIN"], "%d/%m/%Y") < datetime.strptime(values["FECHA_INICIO"], "%d/%m/%Y"):
                errors.append("La fecha de fin es anterior a la fecha de inicio.")
        results.append({"id": index, "fila": row["_fila"], "values": values,
                        "errors": errors, "valid": not errors})
    return {"fields": sorted(tags), "rows": results}


def build_zip(template, prepared, selected):
    valid = [r for r in prepared["rows"] if r["id"] in selected and r["valid"]]
    if not valid:
        raise InputError("Selecciona al menos un registro válido.")
    if len(valid) * len(template) > MAX_OUTPUT_BYTES:
        raise InputError("El lote estimado supera 64 MB. Selecciona menos registros o reduce las imágenes de la plantilla.")
    output = BytesIO()
    report = Workbook()
    sheet = report.active
    sheet.title = "Revisión"
    sheet.append(["Fila", "DNI", "Nombre", "Estado", "Detalle"])
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        for row in prepared["rows"]:
            state = "Generado" if row in valid else ("No seleccionado" if row["valid"] else "Corregir")
            cells = [row["fila"], row["values"].get("DNI", ""), row["values"].get("NOMBRE", ""), state, "; ".join(row["errors"])]
            sheet.append([("'" + c if isinstance(c, str) and c.startswith(("=", "+", "-", "@")) else c) for c in cells])
        total_doc_bytes = 0
        for row in valid:
            doc = Document(BytesIO(template))
            for p in paragraphs(doc):
                replace_paragraph(p, row["values"])
            buffer = BytesIO()
            doc.save(buffer)
            total_doc_bytes += buffer.tell()
            if total_doc_bytes > MAX_OUTPUT_BYTES:
                raise InputError("Los documentos superan 64 MB. Divide el lote en grupos más pequeños.")
            name = re.sub(r"[^\w -]", "", row["values"].get("NOMBRE", "documento"), flags=re.UNICODE).strip()[:70] or "documento"
            archive.writestr(f"documentos/{row['id'] + 1:03d}_{row['values']['DNI']}_{name}.docx", buffer.getvalue())
        buffer = BytesIO()
        report.save(buffer)
        archive.writestr("revision.xlsx", buffer.getvalue())
        archive.writestr("resumen.json", json.dumps({"generados": len(valid), "revisados": len(prepared["rows"])}, ensure_ascii=False))
    return output.getvalue()
